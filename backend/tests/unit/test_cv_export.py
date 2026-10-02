"""Unit tests for app.cv_export - markdown parsing and DOCX/PDF rendering."""
import io
from types import SimpleNamespace

import pytest

from app import cv_export


SAMPLE_CV = """# Jane Doe

## Professional Summary
Engineer with **10 years** experience & a love of <tests>.

## Core Skills
- **Languages:** Python, Go
- **Cloud:** AWS

## Professional Experience
### Staff Engineer • Acme, London @@ Jan 2020 - Present
- Led **platform** team
- Cut costs 30%
### Engineer • Beta Ltd (2015 - 2019)
1. Shipped things
2. Fixed things

## Education
### BSc Computer Science • Uni
"""


def profile(**kw):
    base = dict(full_name="Jane Doe", headline="Staff Engineer", location="London",
                phone="+44 1234", email="jane@example.com", website=None, linkedin=None, github="gh/jane")
    base.update(kw)
    return SimpleNamespace(**base)


class TestSplitBold:
    def test_plain(self):
        assert cv_export._split_bold("hello") == [("hello", False)]

    def test_mixed(self):
        assert cv_export._split_bold("a **b** c **d**") == [
            ("a ", False), ("b", True), (" c ", False), ("d", True)]

    def test_empty(self):
        assert cv_export._split_bold("") == [("", False)]

    def test_unclosed_bold_is_literal(self):
        assert cv_export._split_bold("**oops") == [("**oops", False)]


class TestSplitHeadingDates:
    @pytest.mark.parametrize("text,expected", [
        ("Engineer • Acme @@ 2020 - 2024", ("Engineer • Acme", "2020 - 2024")),
        ("Engineer • Acme @@ Jan 2020 – Present", ("Engineer • Acme", "Jan 2020 – Present")),
        ("Engineer • Acme (2015 - 2019)", ("Engineer • Acme", "2015 - 2019")),
        ("Engineer • Acme | Mar 2018 – Current", ("Engineer • Acme", "Mar 2018 – Current")),
        ("Engineer • Acme — Sept. 2018 - Dec. 2019", ("Engineer • Acme", "Sept. 2018 - Dec. 2019")),
        ("Engineer at Acme", ("Engineer at Acme", None)),
    ])
    def test_cases(self, text, expected):
        assert cv_export._split_heading_dates(text) == expected


class TestParseBlocks:
    def test_structure(self):
        blocks = cv_export._parse_blocks(SAMPLE_CV)
        kinds = [b["kind"] for b in blocks]
        assert kinds == ["h1", "h2", "para", "h2", "bullets", "h2", "entry", "bullets",
                         "entry", "bullets", "h2", "entry"]
        assert blocks[0]["text"] == "Jane Doe"
        assert blocks[4]["items"] == ["**Languages:** Python, Go", "**Cloud:** AWS"]
        assert blocks[6] == {"kind": "entry", "text": "Staff Engineer • Acme, London",
                             "dates": "Jan 2020 - Present"}
        assert blocks[8]["dates"] == "2015 - 2019"
        assert blocks[9]["items"] == ["Shipped things", "Fixed things"]
        assert blocks[11]["dates"] is None

    def test_star_bullets_and_crlf(self):
        blocks = cv_export._parse_blocks("* one\r\n* two\r\n\r\ntext")
        assert blocks == [{"kind": "bullets", "items": ["one", "two"]},
                          {"kind": "para", "text": "text"}]

    def test_empty(self):
        assert cv_export._parse_blocks("") == []


class TestContactLine:
    def test_joins_present_fields_in_order(self):
        assert cv_export._contact_line(profile()) == "London  |  +44 1234  |  jane@example.com  |  gh/jane"

    def test_object_without_attributes(self):
        assert cv_export._contact_line(object()) == ""


class TestBuildDocx:
    def _read(self, data):
        from docx import Document
        return Document(io.BytesIO(data))

    def test_produces_valid_docx_with_header_and_sections(self):
        data = cv_export.build_docx(profile(), SAMPLE_CV)
        assert data[:2] == b"PK"  # zip container
        texts = [p.text for p in self._read(data).paragraphs]
        assert texts[0] == "Jane Doe"
        assert texts[1] == "Staff Engineer"
        assert "jane@example.com" in texts[2]
        # duplicate "# Jane Doe" h1 from the AI body is dropped
        assert texts.count("Jane Doe") == 1
        assert "PROFESSIONAL SUMMARY" in texts
        assert "Staff Engineer • Acme, London\tJan 2020 - Present" in texts
        assert "Languages: Python, Go" in texts

    def test_bold_runs_and_right_tab(self):
        doc = self._read(cv_export.build_docx(profile(), SAMPLE_CV))
        bullet = next(p for p in doc.paragraphs if p.text == "Languages: Python, Go")
        assert [(r.text, bool(r.bold)) for r in bullet.runs] == [("Languages:", True), (" Python, Go", False)]
        entry = next(p for p in doc.paragraphs if p.text.startswith("Staff Engineer • Acme"))
        assert len(entry.paragraph_format.tab_stops) == 1

    def test_missing_profile_fields_fall_back(self):
        doc = self._read(cv_export.build_docx(SimpleNamespace(), "Just a paragraph"))
        assert [p.text for p in doc.paragraphs] == ["Curriculum Vitae", "Just a paragraph"]


class TestBuildPdf:
    def test_produces_pdf(self):
        data = cv_export.build_pdf(profile(), SAMPLE_CV)
        assert data.startswith(b"%PDF")
        assert data.rstrip().endswith(b"%%EOF")

    def test_escapes_html_special_chars(self):
        # "<tests>" would break reportlab's mini-markup if not escaped
        assert cv_export.build_pdf(profile(), "Text with <tag> & stuff").startswith(b"%PDF")

    def test_empty_profile_and_body(self):
        assert cv_export.build_pdf(SimpleNamespace(), "").startswith(b"%PDF")
